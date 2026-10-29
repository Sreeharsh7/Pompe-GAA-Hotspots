"""Stage 9: assemble the Supplementary Tables workbook (one titled sheet per table) and CSV copies.

Outputs
  results/supplementary/Supplementary_Tables.xlsx
  results/supplementary/csv/Table_S<n>_<slug>.csv
"""

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from gaa_hotspot.common import AA, MODEL_NAMES, MODEL_ORDER, config, load_values, path, read_processed, table
from gaa_hotspot.ml import GRIDS

cfg = config()
V = load_values()
out = path("supplementary")
(out / "csv").mkdir(exist_ok=True)

feat = read_processed("residue_features.csv")
groups = read_processed("residue_network_groups.csv")
votes = pd.read_csv(table("residue_votes.csv"))
novel = pd.read_csv(table("novel_hotspots.csv"))
sss = read_processed("sss_all_substitutions.csv")

grid_rows = []
for m in MODEL_ORDER:
    g = GRIDS[m]
    for i, sub in enumerate(g if isinstance(g, list) else [g]):
        for k, v in sub.items():
            grid_rows.append({"Model": MODEL_NAMES[m], "Sub-grid": i + 1, "Hyperparameter": k, "Values searched": ", ".join(map(str, v))})
grid = pd.DataFrame(grid_rows)

params = pd.read_csv(table("selected_hyperparameters.csv"))
params["model"] = params.model.map(MODEL_NAMES)
params["repeat"] += 1
params["fold"] += 1

nov_mat = sss[sss.residue.isin(novel.residue)].pivot(index="mt", columns="residue", values="sss").reindex(AA)
nov_mat.columns = [f"{feat.set_index('residue').wt[c]}{c}" for c in nov_mat.columns]
nov_mat.index.name = "Substituted amino acid"

S = [
    ("Residue-level features and mutation labels",
     "All 844 structurally resolved GAA residues (PDB 5NN3, Rosetta-relaxed). RIN centralities were computed with python-igraph from the RINerator network; "
     "structural attributes from UCSF Chimera; conservation is the normalised ConSurf-DB score (lower = more conserved). mutation_observed = 1 for positions with a reported missense mutation. "
     "log_betweenness = log10(1 + betweenness).",
     feat),
    ("Univariate comparison of features between mutation positions and other residues",
     "Two-sided permutation test of the difference in means (10,000 label permutations; minimum attainable P = 1e-4), Benjamini-Hochberg q-values across features, Cohen's d and Mann-Whitney U P-values.",
     pd.read_csv(table("feature_statistics.csv"))),
    ("Spearman correlation between RIN centrality measures", "Pairwise Spearman rank correlation across 844 residues.",
     pd.read_csv(table("centrality_correlation.csv"))),
    ("Topological criticality groups and supercritical hubs",
     f"K-means binning ({cfg['criticality']['kbins']} bins) of degree and log betweenness; HDHB = top bin on both, LDHB = bottom degree and top betweenness bin, LDLB = bottom bin on both. "
     "Supercritical hubs are the 10 residues closest (Euclidean, unscaled axes) to the ideal point defined by the maxima of degree, log betweenness and closeness; "
     "supercritical_z repeats this after z-scoring each axis (sensitivity analysis).",
     groups),
    ("Hyperparameter search spaces", f"Inner 5-fold grid search optimising ROC AUC. Random forest: {cfg['ml']['rf_random_search_iter']} configurations sampled at random from the full grid; all other grids exhaustive.", grid),
    ("Hyperparameters selected in each outer fold", "Inner-CV ROC AUC of the selected configuration and ROC AUC on the held-out outer fold, for 5 repeats x 10 outer folds.", params),
    ("Classifier performance from repeated nested cross-validation",
     "pooled_oof_auc: ROC AUC of the mean out-of-fold probability across 5 repeats, with DeLong 95% CI. per_repeat_auc: pooled AUC within each repeat (mean, SD). "
     f"fold_auc: AUC per outer fold (mean, SD). Threshold metrics at probability >= {cfg['ml']['vote_threshold']}.",
     pd.read_csv(table("model_performance.csv"))),
    ("Pairwise DeLong tests between classifiers", "Best model (highest pooled OOF AUC) against each other model; p_holm = Holm-adjusted P.", pd.read_csv(table("delong_comparisons.csv"))),
    ("Per-residue out-of-fold probabilities, votes and classification",
     f"A model votes for a residue when its mean OOF probability reaches the model's {int(100 * cfg['ml']['vote_specificity'])} % specificity threshold (Table S22); consensus = at least {cfg['ml']['consensus_min_votes']} of 5 votes. "
     "consensus_repeat_fraction = fraction of the 5 independent CV repeats in which the residue reached consensus. Distances are minimum heavy-atom distances to the nearest active-site residue.",
     votes),
    ("Novel consensus hotspots", "Consensus residues without a reported mutation.", novel),
    ("Domain distribution of novel consensus hotspots", "expected_by_size = number expected if hotspots were distributed in proportion to domain size; two-sided binomial test.", pd.read_csv(table("hotspot_domain_breakdown.csv"))),
    ("Known mutation positions not prioritised by any model", "Medians for known positions with zero votes versus known positions reaching consensus; Mann-Whitney U test.", pd.read_csv(table("discordant_positions.csv"))),
    ("Permutation feature importance", "Decrease in held-out ROC AUC when a feature is permuted (20 permutations), averaged over the 10 outer folds of the first CV repeat.", pd.read_csv(table("permutation_importance.csv"))),
    ("SHAP feature importance", "Mean |SHAP| for logistic regression (linear explainer), random forest and XGBoost (tree explainer) refitted on all residues; mean_normalised averages values scaled to each model's maximum.", pd.read_csv(table("shap_importance.csv"))),
    ("Physicochemical distance index (DI)", f"Euclidean distance between amino acids in the space of the first {V['sss']['n_components']} principal components ({100 * V['sss']['variance_retained']:.1f}% variance) of {V['sss']['n_properties']} standardised AAindex properties (SeqinR).", pd.read_csv(table("di_matrix.csv"))),
    ("Substitution Severity Score at the novel consensus hotspots",
     "SSS = (EVOscore/3 clipped to [0,1] + DI/max DI) / 2; rows are substituted amino acids, columns hotspot residues (wild type + number). Wild-type cells are empty.", nov_mat.reset_index()),
    ("EVOscore, DI and SSS for every possible substitution",
     f"EVOscore = -log10(f_variant / f_wildtype) from ConSurf alignment frequencies; residues absent from the alignment take f = {cfg['sss']['min_frequency']}, '<1%' entries f = {cfg['sss']['below_one_percent']}.", sss),
    ("ClinVar GAA missense variants with predictor scores",
     "ClinVar snapshot distributed with the EVE GAA model. ESM-1v is reported as the negated log-likelihood ratio (higher = more damaging). all_four_scores marks variants scored by SSS, AlphaMissense, ESM-1v and EVE.",
     pd.read_csv(table("benchmark_variants.csv"))),
    ("Benchmark of SSS against variant effect predictors",
     f"P/LP versus B/LB ClinVar variants scored by all four predictors (n = {V['benchmark']['benchmark_n']}: {V['benchmark']['benchmark_pathogenic']} P/LP, {V['benchmark']['benchmark_benign']} B/LB). "
     "DeLong and stratified bootstrap (2,000 resamples) 95% CIs; Spearman correlation between score and class; Mann-Whitney U test. EVOscore and DI are the two SSS components.",
     pd.read_csv(table("benchmark_performance.csv"))),
    ("Paired DeLong tests of SSS against each predictor", "Same benchmark variants as Table S19.", pd.read_csv(table("benchmark_delong_vs_sss.csv"))),
    ("SSS classification of benchmark variants at fixed and data-derived thresholds", "The Youden-optimal threshold is derived from the same variants and is therefore optimistic.", pd.read_csv(table("benchmark_threshold.csv"))),
]
S.append(("Model-specific vote thresholds", f"Probability at which each model reaches {int(100 * cfg['ml']['vote_specificity'])} % specificity on residues without a reported mutation (mean OOF probability).", pd.read_csv(table("vote_thresholds.csv"))))
S.append(("Sensitivity of the consensus hotspot set to the vote rule", "Consensus = at least 3 of 5 model votes under each rule.", pd.read_csv(table("vote_threshold_sensitivity.csv"))))
if table("feature_ablation.csv").exists():
    S.append(("Feature-group ablation", "Nested CV (2 repeats x 10 folds) on feature subsets; DeLong P against the full 14-feature model.", pd.read_csv(table("feature_ablation.csv"))))
if table("legacy811_model_performance.csv").exists():
    S.append(("Sensitivity analysis: original 811-residue subset",
              "Repeated nested CV restricted to the 811 residues used in earlier drafts (zero-betweenness residues and 8 additional residues removed; betweenness as log10).",
              pd.read_csv(table("legacy811_model_performance.csv"))))

wb = Workbook()
toc = wb.active
toc.title = "Contents"
toc["A1"] = "Supplementary Tables — Uncovering novel pathogenic hotspots in the Pompe disease target using graph-based features and ensemble machine learning"
toc["A1"].font = Font(bold=True, size=12)
toc.append([])
toc.append(["Table", "Title", "Rows"])
for c in toc[3]:
    c.font = Font(bold=True)
head_fill = PatternFill("solid", fgColor="E8EEF4")
for i, (title, legend, df) in enumerate(S, start=1):
    name = f"Table S{i}"
    slug = "_".join(title.split()[:6]).replace(",", "").replace("(", "").replace(")", "").replace("/", "-")
    df.to_csv(out / "csv" / f"Table_S{i}_{slug}.csv", index=False)
    ws = wb.create_sheet(f"S{i}")
    ws["A1"] = f"Supplementary {name}. {title}"
    ws["A1"].font = Font(bold=True, size=11)
    ws["A2"] = legend
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(8, min(len(df.columns), 14)))
    ws.row_dimensions[2].height = 15 * max(2, len(legend) // 140 + 1)
    ws.append([])
    ws.append([str(c) for c in df.columns])
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else (json.dumps(v) if isinstance(v, (dict, list)) else v) for v in row])
    for j, col in enumerate(df.columns, start=1):
        width = max(len(str(col)), *(len(f"{v:.4g}" if isinstance(v, float) else str(v)) for v in df[col].head(200))) if len(df) else len(str(col))
        ws.column_dimensions[get_column_letter(j)].width = min(max(9, width + 2), 48)
        if pd.api.types.is_float_dtype(df[col]):
            for cell in ws.iter_rows(min_row=5, min_col=j, max_col=j):
                cell[0].number_format = "0.0000"
    ws.freeze_panes = "A5"
    toc.append([name, title, len(df)])
toc.column_dimensions["A"].width = 10
toc.column_dimensions["B"].width = 95
wb.save(out / "Supplementary_Tables.xlsx")
print(f"{len(S)} supplementary tables written to {out}")
