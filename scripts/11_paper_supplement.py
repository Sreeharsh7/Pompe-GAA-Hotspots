"""Stage 11: the Supplementary Information exactly as submitted with the paper.

Writes results/supplementary/paper/Supplementary_Information.pdf (Figs. S1-S5 with legends) and
Supplementary_Tables.xlsx (Tables S1-S4). Numbering follows the order of citation in the paper;
the repository's own FigS*/Table S* files are a larger set with a different numbering (see README).
"""

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import numpy as np
import pandas as pd
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, PageBreak, NextPageTemplate, PageTemplate, Frame
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[1]
OUT = REPO / "results" / "supplementary" / "paper"
OUT.mkdir(parents=True, exist_ok=True)
TMP = Path(tempfile.mkdtemp())
SUPP = REPO / "results/figures/supplementary"; TAB = REPO / "results/tables"
TITLE = "Uncovering novel pathogenic hotspots in the Pompe disease target using graph-based features and ensemble machine learning"
AUTHORS = "Sreeharsh K, Manjari Kiran, Mohd. Akif"
Image.MAX_IMAGE_PIXELS = None

def shrink(src, dst, max_w=2600):
    im = Image.open(src).convert("RGB")
    if im.width > max_w:
        im = im.resize((max_w, round(im.height * max_w / im.width)), Image.LANCZOS)
    im.save(dst, dpi=(300, 300)); return dst

# Fig. S4 composite: permutation importance + three SHAP panels
def composite():
    files = ["FigS5a_permutation_importance.png", "FigS5b_SHAP_RF.png", "FigS5c_SHAP_LR.png", "FigS5d_SHAP_XGB.png"]
    ims = [Image.open(SUPP / f).convert("RGB") for f in files]
    w = 2200; ims = [i.resize((w, round(i.height * w / i.width)), Image.LANCZOS) for i in ims]
    h = max(i.height for i in ims); pad = 140
    canvas = Image.new("RGB", (2 * w + pad, 2 * h + pad), "white")
    try: font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 110)
    except OSError: font = ImageFont.load_default()
    d = ImageDraw.Draw(canvas)
    for k, (im, lab) in enumerate(zip(ims, "abcd")):
        x = (k % 2) * (w + pad); y = (k // 2) * (h + pad)
        canvas.paste(im, (x + pad // 2 if k % 2 == 0 else x, y + pad))
        d.text((x + 10 if k % 2 == 0 else x - 30, y + 10), lab, fill="black", font=font)
    p = TMP / "FigS4_feature_importance.png"; canvas.save(p, dpi=(300, 300)); return p

FIGS = [
    ("FigS2_model_metrics.png", False, "Performance of the five classifiers in repeated nested cross-validation",
     "(a) Accuracy, (b) F1 score, (c) ROC AUC and (d) balanced accuracy of Random Forest, Support Vector Machine, Logistic Regression, "
     "k-nearest neighbors and XGBoost at a probability threshold of 0.5. Bars show the mean over the 50 outer test folds "
     "(5 repeats × 10 folds) and error bars the standard deviation. Fold-averaged AUCs are therefore slightly lower than the pooled "
     "out-of-fold AUCs reported in Fig. 4a."),
    ("FigS3_confusion_matrices.png", True, "Confusion matrices of the five classifiers",
     "Out-of-fold predictions for all 844 residues (214 mutation positions, 630 other residues). Top row: probability threshold of 0.5. "
     "Bottom row: model-specific threshold at 90% specificity, which was used for consensus voting. Each cell gives the number and "
     "percentage of residues; panel titles give the threshold, sensitivity and specificity. LR, Logistic Regression; RF, Random Forest; "
     "kNN, k-nearest neighbors; SVM, Support Vector Machine; XGB, XGBoost."),
    ("FigS4_cv_ROC_mean_sd.png", False, "Cross-validated ROC curves",
     "Mean ROC curve of each classifier over the 50 outer folds of repeated nested cross-validation; shaded bands show ±1 standard "
     "deviation. Legend values are the mean ± standard deviation of the fold AUCs. The dashed line represents random prediction."),
    (None, False, "Feature contributions to classifier predictions",
     "(a) Permutation importance: decrease in held-out ROC AUC when each feature is randomly permuted in the outer test folds, shown for "
     "each classifier. (b–d) SHAP summary plots for Random Forest (b), Logistic Regression (c) and XGBoost (d) refitted on all residues "
     "with the selected hyperparameters; each point is one residue, coloured by its feature value. Conserv_score, ConSurf conservation "
     "score; kdhydro, Kyte–Doolittle hydrophobicity; areaSAS and areaSES, solvent-accessible and solvent-excluded surface area; "
     "relSESA, relative solvent-excluded surface area; burt, Burt's constraint; k_core, K-core; page_rank, PageRank-like score; "
     "phi and psi, backbone dihedral angles."),
    ("FigS7_hotspot_stability.png", True, "Stability of the novel consensus hotspots across cross-validation repeats",
     "Fraction of the five cross-validation repeats in which each of the 52 novel consensus hotspots reached consensus (at least three of "
     "five classifier votes). Bars are coloured by the number of votes in the pooled analysis. Thirty-three positions reached consensus "
     "in all five repeats."),
]

ss = getSampleStyleSheet()
cap = ParagraphStyle("cap", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=13, alignment=TA_JUSTIFY)
h1 = ParagraphStyle("h1", parent=ss["Title"], fontName="Helvetica-Bold", fontSize=15, leading=20, alignment=TA_CENTER)
nm = ParagraphStyle("nm", parent=ss["Normal"], fontName="Helvetica", fontSize=11, leading=15, alignment=TA_CENTER)
li = ParagraphStyle("li", parent=ss["Normal"], fontName="Helvetica", fontSize=10, leading=15)

M = 1.6 * cm
pw, ph = A4; lw, lh = landscape(A4)
doc = SimpleDocTemplate(str(OUT / "Supplementary_Information.pdf"), pagesize=A4, title="Supplementary Information", author=AUTHORS)
doc.addPageTemplates([
    PageTemplate(id="P", frames=[Frame(M, M, pw - 2 * M, ph - 2 * M)], pagesize=A4),
    PageTemplate(id="L", frames=[Frame(M, M, lw - 2 * M, lh - 2 * M)], pagesize=landscape(A4)),
])
story = [Spacer(1, 3 * cm), Paragraph("Supplementary Information", h1), Spacer(1, 0.6 * cm), Paragraph(TITLE, nm),
         Spacer(1, 0.4 * cm), Paragraph(AUTHORS, nm), Spacer(1, 1.5 * cm), Paragraph("<b>Contents</b>", li), Spacer(1, 0.2 * cm)]
for k, f in enumerate(FIGS, 1):
    story.append(Paragraph(f"Supplementary Figure S{k}. {f[2]}", li))
story.append(Paragraph("Supplementary Tables S1–S4 are provided as a separate Excel workbook (Supplementary_Tables.xlsx).", li))

for k, (fname, land, title, legend) in enumerate(FIGS, 1):
    src = composite() if fname is None else SUPP / fname
    img = shrink(src, TMP / f"S{k}.png", 3000 if land else 2400)
    story += [NextPageTemplate("L" if land else "P"), PageBreak()]
    fw = (lw if land else pw) - 2 * M; fh = (lh if land else ph) - 2 * M - (4.2 if land else 5.5) * cm
    iw, ih = Image.open(img).size; s = min(fw / iw, fh / ih)
    story += [RLImage(str(img), width=iw * s, height=ih * s), Spacer(1, 0.4 * cm),
              Paragraph(f"<b>Supplementary Figure S{k} | {title}.</b> {legend}", cap)]
doc.build(story)

# ---------------- Excel ----------------
FEAT = {"degree": "Degree", "betweenness": "Betweenness", "log_betweenness": "Betweenness (log10(1 + x))", "closeness": "Closeness", "burt": "Burt's constraint",
        "authority": "Authority", "page_rank": "PageRank-like score", "k_core": "K-core", "kdhydro": "Hydrophobicity (Kyte–Doolittle)",
        "areaSAS": "Solvent-accessible surface area (Å²)", "areaSES": "Solvent-excluded surface area (Å²)", "relSESA": "Relative SESA",
        "phi": "Backbone φ (°)", "psi": "Backbone ψ (°)", "conservation": "ConSurf conservation score"}
fs = pd.read_csv(TAB / "feature_statistics.csv"); fs = fs[fs.feature != "betweenness"]
t1 = pd.DataFrame({"Feature": fs.feature.map(FEAT),
    "Mean (mutation positions, n = 214)": fs.mean_mutation_positions, "Mean (other residues, n = 630)": fs.mean_other,
    "Median (mutation positions)": fs.median_mutation_positions, "Median (other residues)": fs.median_other,
    "Cohen's d": fs.cohens_d, "Permutation P": fs.permutation_p, "BH-adjusted q": fs.permutation_q_bh, "Mann–Whitney P": fs.mann_whitney_p})
mp = pd.read_csv(TAB / "model_performance.csv")
NAMES = {"lr": "Logistic Regression", "rf": "Random Forest", "svm": "Support Vector Machine", "xgb": "XGBoost", "knn": "k-nearest neighbors"}
t2 = pd.DataFrame({"Classifier": mp.model.map(NAMES), "Pooled out-of-fold ROC AUC": mp.pooled_oof_auc,
    "95% CI (lower)": mp.auc_ci_low, "95% CI (upper)": mp.auc_ci_high, "Fold ROC AUC, mean": mp.fold_auc_mean, "Fold ROC AUC, SD": mp.fold_auc_sd,
    "Sensitivity": mp.sensitivity, "Specificity": mp.specificity, "Precision": mp.precision, "F1 score": mp.f1,
    "Balanced accuracy": mp.balanced_accuracy, "MCC": mp.mcc}).sort_values("Pooled out-of-fold ROC AUC", ascending=False)
nh = pd.read_csv(TAB / "novel_hotspots.csv").sort_values("residue")
DOM = {"Trefoil type-P": "Trefoil type-P (81–136)", "N-terminal beta-sheet": "N-terminal β-sheet (137–357)",
       "Catalytic (beta/alpha)8 barrel": "Catalytic (β/α)8 barrel (358–720)", "Proximal beta-sheet": "Proximal β-sheet (721–820)",
       "Distal beta-sheet": "Distal β-sheet (821–952)"}
yn = lambda s: s.map({1: "Yes", 0: "No", True: "Yes", False: "No"})
t3 = pd.DataFrame({"Residue": nh.residue_label, "Domain": nh.domain.map(DOM), "Votes (of 5)": nh.votes,
    "LR": yn(nh.vote_lr), "RF": yn(nh.vote_rf), "SVM": yn(nh.vote_svm), "XGB": yn(nh.vote_xgb), "kNN": yn(nh.vote_knn),
    "Mean out-of-fold probability": nh.mean_probability, "Fraction of CV repeats reaching consensus": nh.consensus_repeat_fraction,
    "Criticality group": nh.criticality.replace("Other", "—"), "Supercritical hub": yn(nh.supercritical),
    "SASA (Å²)": nh.areaSAS, "ConSurf conservation score": nh.conservation, "Active-site residue": yn(nh.is_active_site_residue),
    "Nearest active-site residue": nh.nearest_active_site_residue, "Distance to nearest active-site residue (Å)": nh.nearest_active_site_distance_A})
bv = pd.read_csv(TAB / "benchmark_variants.csv")
bench = bv["class"].isin(["Pathogenic", "Benign"]) & bv.all_four_scores
CLS = {"Pathogenic": "Pathogenic/likely pathogenic", "Benign": "Benign/likely benign", "VUS": "Uncertain significance", "Conflicting": "Conflicting"}
t4 = pd.DataFrame({"Variant": "p." + bv.wt + bv.residue.astype(str) + bv.mt, "ClinVar classification": bv.clinvar,
    "Class used": bv["class"].map(CLS).fillna(bv["class"]), "In benchmark set (n = 102)": yn(bench),
    "SSS": bv.SSS, "EVOscore": bv.EVOscore, "DI": bv.DI, "AlphaMissense": bv.AlphaMissense, "ESM-1v (−LLR)": bv["ESM-1v"], "EVE": bv.EVE})

TABLES = [
    ("Table S1", "Univariate comparison of residue features between reported mutation positions and other residues",
     "Two-sided permutation test of the difference in means (10,000 label permutations; the smallest attainable P is 1.0 × 10⁻⁴), "
     "with Benjamini–Hochberg adjustment (q). Cohen's d is positive when mutation positions have higher values. Mann–Whitney P is given "
     "for reference. φ and ψ are missing for three and two residues, respectively.", t1),
    ("Table S2", "Performance of the five classifiers in repeated nested cross-validation",
     "5 repeats × stratified 10-fold outer cross-validation with 5-fold inner grid search. Pooled ROC AUC uses the mean out-of-fold "
     "probability of each residue across repeats, with DeLong 95% CI. Fold AUC is the mean and SD over the 50 outer test folds. "
     "Sensitivity, specificity, precision, F1, balanced accuracy and Matthews correlation coefficient (MCC) are at a probability threshold of 0.5.", t2),
    ("Table S3", "Novel consensus hotspot positions (n = 52)",
     "Residues without a reported mutation that received votes from at least three of the five classifiers (each classifier votes when "
     "the mean out-of-fold probability exceeds its threshold at 90% specificity). Distances are minimum heavy-atom distances in the "
     "relaxed 5NN3 structure.", t3),
    ("Table S4", "ClinVar GAA missense variants and predictor scores",
     "All 574 ClinVar GAA missense variants. The benchmark set comprises the 102 pathogenic/likely pathogenic or benign/likely benign "
     "variants scored by all four predictors. ESM-1v log-likelihood ratios are negated so that higher values indicate more damaging "
     "substitutions for every predictor. Empty cells: no score available (e.g. residues outside the resolved structure for SSS).", t4),
]
wb = Workbook(); ws = wb.active; ws.title = "Contents"
bold = Font(bold=True); hdr = PatternFill("solid", fgColor="DDE4EE"); thin = Side(style="thin", color="999999")
ws["A1"] = "Supplementary Tables"; ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = TITLE; ws["A3"] = AUTHORS; ws["A2"].font = Font(italic=True)
ws["A5"], ws["B5"] = "Table", "Title"; ws["A5"].font = ws["B5"].font = bold
for r, (code, title, _, df) in enumerate(TABLES, 6):
    ws.cell(r, 1, code).hyperlink = f"#'{code}'!A1"; ws.cell(r, 1).font = Font(color="1F4E9D", underline="single"); ws.cell(r, 2, title)
ws.column_dimensions["A"].width = 12; ws.column_dimensions["B"].width = 100
for code, title, note, df in TABLES:
    s = wb.create_sheet(code)
    s["A1"] = f"Supplementary {code}. {title}"; s["A1"].font = Font(bold=True, size=12)
    s["A2"] = note; s["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    s.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(6, len(df.columns))); s.row_dimensions[2].height = 48
    for j, c in enumerate(df.columns, 1):
        cell = s.cell(4, j, c); cell.font = bold; cell.fill = hdr
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center"); cell.border = Border(bottom=thin, top=thin)
    s.row_dimensions[4].height = 45
    for i, row in enumerate(df.itertuples(index=False), 5):
        for j, v in enumerate(row, 1):
            if isinstance(v, float) and np.isnan(v): v = None
            elif isinstance(v, (np.floating,)): v = float(v)
            elif isinstance(v, (np.integer,)): v = int(v)
            c = s.cell(i, j, v)
            if isinstance(v, float):
                col = df.columns[j - 1]
                c.number_format = "0.00E+00" if ("P" in col.split() or col.endswith(" q") or "q" == col[-1:]) and v < 0.001 else "0.000"
    for j, c in enumerate(df.columns, 1):
        s.column_dimensions[get_column_letter(j)].width = max(11, min(32, len(c) * 0.55 + 6)) if j > 1 else max(14, df.iloc[:, 0].astype(str).str.len().max() + 3)
    s.freeze_panes = "B5"
wb.save(OUT / "Supplementary_Tables.xlsx")
print("Supplementary Information written to", OUT, [(t[0], len(t[3])) for t in TABLES])
